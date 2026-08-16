#!/usr/bin/env python3
"""Attendance Wave 2B — CSV/XLSX/SFTP fallback through the canonical punch contract."""

from __future__ import annotations

import csv
import hashlib
import io
from pathlib import Path
from typing import Any, Iterable

from attendance_capture_contract import (
    CONNECTOR_VERSION,
    CanonicalPunch,
    as_kuwait,
    csv_source_event_id,
    normalize_capture_method,
    sanitize_vendor_payload,
    sftp_source_event_id,
    stable_hash,
)

DEFAULT_COLUMN_MAP = {
    "device_user_id": ("device_user_id", "emp_code", "employee_id", "badge", "user_id", "pin"),
    "punched_at": ("punched_at", "punch_time", "timestamp", "time", "datetime", "check_time"),
    "direction": ("direction", "punch_state", "status", "inout", "type"),
    "device_id": ("device_id", "terminal_sn", "device_sn", "sn", "terminal"),
    "capture_method": ("capture_method", "verify_type", "method", "auth_type"),
    "employee_key": ("employee_key",),
}

DIRECTION_ALIASES = {
    "0": "in",
    "1": "out",
    "i": "in",
    "o": "out",
    "in": "in",
    "out": "out",
    "check_in": "in",
    "check_out": "out",
    "checkin": "in",
    "checkout": "out",
    "break_start": "break_start",
    "break_end": "break_end",
    "breakstart": "break_start",
    "breakend": "break_end",
}


def file_sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _pick(row: dict[str, Any], aliases: Iterable[str]) -> Any:
    lower = {str(k).strip().lower(): v for k, v in row.items()}
    for a in aliases:
        if a.lower() in lower and lower[a.lower()] not in (None, ""):
            return lower[a.lower()]
    return None


def _normalize_direction(value: Any) -> str | None:
    key = str(value or "").strip().lower()
    return DIRECTION_ALIASES.get(key)


def _rows_from_csv_bytes(data: bytes) -> list[dict[str, Any]]:
    text = data.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text))
    return [dict(r) for r in reader]


def _rows_from_xlsx_bytes(data: bytes) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    wb = load_workbook(filename=io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    headers = [str(h).strip() if h is not None else f"col{i}" for i, h in enumerate(next(rows_iter))]
    out: list[dict[str, Any]] = []
    for row in rows_iter:
        item = {headers[i]: row[i] for i in range(len(headers))}
        if all(v is None or str(v).strip() == "" for v in item.values()):
            continue
        out.append(item)
    return out


def load_tabular_rows(path: str | Path | None = None, *, data: bytes | None = None, filename: str | None = None) -> tuple[list[dict[str, Any]], str, str]:
    if path is not None:
        p = Path(path)
        data = p.read_bytes()
        filename = filename or p.name
    if data is None:
        raise ValueError("missing_file_bytes")
    name = (filename or "upload.csv").lower()
    fsha = file_sha256(data)
    if name.endswith((".xlsx", ".xlsm")):
        rows = _rows_from_xlsx_bytes(data)
    else:
        rows = _rows_from_csv_bytes(data)
    return rows, fsha, name


def row_to_canonical(
    row: dict[str, Any],
    *,
    company_code: str,
    connector_id: str,
    file_sha: str,
    source: str = "csv",
    column_map: dict[str, tuple[str, ...]] | None = None,
) -> dict[str, Any]:
    san = sanitize_vendor_payload(row)
    if not san.ok:
        return {"ok": False, "error": san.reason, "forbidden_keys": san.forbidden_keys, "privacy_hard_fail": True}
    clean = san.payload or {}
    cmap = column_map or DEFAULT_COLUMN_MAP
    device_user_id = str(_pick(clean, cmap["device_user_id"]) or "").strip()
    punched_at = as_kuwait(_pick(clean, cmap["punched_at"]))
    direction = _normalize_direction(_pick(clean, cmap["direction"]))
    device_id = _pick(clean, cmap["device_id"])
    device_id_s = str(device_id).strip() if device_id not in (None, "") else None
    capture_method = normalize_capture_method(_pick(clean, cmap["capture_method"]) or "file_import")
    if capture_method == "unknown":
        capture_method = "file_import"
    employee_key = _pick(clean, cmap.get("employee_key", ("employee_key",)))
    employee_key_s = str(employee_key).strip() if employee_key not in (None, "") else None

    row_sha = stable_hash(
        device_user_id,
        punched_at.isoformat() if punched_at else "",
        direction or "",
        device_id_s or "",
        capture_method,
        json_safe_row(clean),
    )
    if source == "sftp":
        source_event_id = sftp_source_event_id(file_sha, row_sha)
    else:
        source_event_id = csv_source_event_id(file_sha, row_sha)

    if not device_user_id:
        return {"ok": False, "error": "missing_device_user_id", "source_event_id": source_event_id, "quarantine": True}
    if punched_at is None:
        return {"ok": False, "error": "invalid_punched_at", "source_event_id": source_event_id, "quarantine": True}
    if direction is None:
        return {
            "ok": False,
            "error": "unmapped_direction",
            "source_event_id": source_event_id,
            "quarantine": True,
            "device_user_id": device_user_id,
            "device_id": device_id_s,
        }

    punch = CanonicalPunch(
        company_code=company_code.upper(),
        source=source if source in {"csv", "sftp", "file_import"} else "csv",
        source_event_id=source_event_id,
        device_user_id=device_user_id,
        punched_at=punched_at,
        direction=direction,
        capture_method=capture_method,
        employee_key=employee_key_s,
        device_id=device_id_s,
        connector_id=connector_id,
        connector_version=CONNECTOR_VERSION,
        raw_ref={
            "file_sha256": file_sha,
            "row_sha256": row_sha,
        },
        metadata={"capture_method": capture_method, "import_kind": source},
    )
    return {"ok": True, "punch": punch}


def json_safe_row(row: dict[str, Any]) -> str:
    parts = []
    for k in sorted(row.keys(), key=lambda x: str(x)):
        parts.append(f"{k}={row.get(k)}")
    return "&".join(parts)


def parse_file_to_canonical(
    *,
    company_code: str,
    connector_id: str,
    path: str | Path | None = None,
    data: bytes | None = None,
    filename: str | None = None,
    source: str = "csv",
    column_map: dict[str, tuple[str, ...]] | None = None,
) -> dict[str, Any]:
    rows, fsha, name = load_tabular_rows(path, data=data, filename=filename)
    accepted: list[CanonicalPunch] = []
    quarantined: list[dict[str, Any]] = []
    privacy_fails: list[dict[str, Any]] = []
    seen: set[str] = set()
    for row in rows:
        mapped = row_to_canonical(
            row,
            company_code=company_code,
            connector_id=connector_id,
            file_sha=fsha,
            source=source,
            column_map=column_map,
        )
        if mapped.get("privacy_hard_fail"):
            privacy_fails.append(mapped)
            continue
        if not mapped.get("ok"):
            quarantined.append(mapped)
            continue
        punch: CanonicalPunch = mapped["punch"]
        if punch.source_event_id in seen:
            continue
        seen.add(punch.source_event_id)
        accepted.append(punch)
    return {
        "ok": True,
        "filename": name,
        "file_sha256": fsha,
        "punches": accepted,
        "quarantined": quarantined,
        "privacy_fails": privacy_fails,
        "row_count": len(rows),
    }


class SftpFileSource:
    """Thin wrapper: caller supplies listed local paths synced from SFTP drop."""

    def __init__(self, drop_dir: str | Path) -> None:
        self.drop_dir = Path(drop_dir)

    def list_files(self) -> list[Path]:
        if not self.drop_dir.is_dir():
            return []
        out = []
        for p in sorted(self.drop_dir.iterdir()):
            if p.suffix.lower() in {".csv", ".xlsx", ".xlsm", ".txt"} and p.is_file():
                out.append(p)
        return out

    def ingest_all(self, *, company_code: str, connector_id: str) -> dict[str, Any]:
        batches = []
        for path in self.list_files():
            batches.append(
                parse_file_to_canonical(
                    company_code=company_code,
                    connector_id=connector_id,
                    path=path,
                    source="sftp",
                )
            )
        return {"ok": True, "batches": batches, "files": len(batches)}
